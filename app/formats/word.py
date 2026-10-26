# -*- coding: utf-8 -*-
"""
app/formats/word.py

將 ScoreDocument 匯出為與 Score Studio 編輯頁視覺一致的 Word (.docx) 檔案。
特點：
- 窄邊界（左右 1.0 cm，總寬 19.0 cm）搭配每行 4 小節依內容長度動態分配欄寬
- 簡譜音符支援高低八度上下點（\u0307 / \u0323）、減時線底線（單底線 _ / 雙底線 =）、連音線 ⌒ / 圓滑線 ⌣
- 音符與歌詞強制單行不折行（適配字級與空白壓縮），避免跑到下一行
- 標頭與編輯區一致：左方「1=C  4/4  ♩=80」、右方「作詞  作曲」
"""

from __future__ import annotations
import io
import re
from typing import Any, Dict, List, Tuple

import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_UNDERLINE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SECTION_LABELS = {
    "intro": "前奏",
    "verse": "主歌",
    "prechorus": "導歌",
    "chorus": "副歌",
    "bridge": "橋段",
    "interlude": "間奏",
    "solo": "間奏/Solo",
    "outro": "尾奏",
    "coda": "尾聲",
}


def _set_east_asia_font(run, font_name: str = "微軟正黑體", ascii_font: str = "Arial"):
    run.font.name = ascii_font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), font_name)


def _set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge, spec in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        existing = tcBorders.find(qn(f"w:{edge}"))
        if existing is not None:
            tcBorders.remove(existing)
        el = OxmlElement(f"w:{edge}")
        if spec:
            el.set(qn("w:val"), spec.get("val", "single"))
            el.set(qn("w:sz"), str(spec.get("sz", 6)))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), spec.get("color", "444444"))
        else:
            el.set(qn("w:val"), "nil")
        tcBorders.append(el)


def _set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.find(qn("w:tcMar"))
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for side, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        existing = tcMar.find(qn(f"w:{side}"))
        if existing is not None:
            tcMar.remove(existing)
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)


def _set_cell_width_cm(cell, width_cm: float):
    cell.width = Cm(width_cm)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(round(width_cm * 567))))
    tcW.set(qn("w:type"), "dxa")


def _extract_measure_texts(m: Dict[str, Any]) -> Tuple[str, str, List[str]]:
    chord_text = str(m.get("chordText") or "").strip()
    if not chord_text and isinstance(m.get("chords"), list):
        ch_list = []
        for c in m.get("chords", []):
            if isinstance(c, dict):
                sym = str(c.get("symbol") or c.get("chord") or c.get("text") or "").strip()
                if sym:
                    ch_list.append(sym)
            elif isinstance(c, str) and c.strip():
                ch_list.append(c.strip())
        chord_text = "  ".join(ch_list)

    note_text = str(m.get("noteText") or "").rstrip()
    if not note_text and isinstance(m.get("notes"), list):
        n_list = []
        for n in m.get("notes", []):
            if isinstance(n, dict):
                tok = str(n.get("token") or n.get("pitch") or n.get("text") or "").strip()
                if tok:
                    n_list.append(tok)
            elif isinstance(n, str) and n.strip():
                n_list.append(n.strip())
        note_text = " ".join(n_list)

    lyric_lines: List[str] = []
    if isinstance(m.get("lyricLines"), list) and m["lyricLines"]:
        lyric_lines = [str(line or "").rstrip() for line in m["lyricLines"]]
    else:
        lt = str(m.get("lyricText") or "").rstrip()
        if not lt and isinstance(m.get("lyrics"), list):
            l_list = []
            for l_item in m.get("lyrics", []):
                if isinstance(l_item, dict):
                    txt = str(l_item.get("text") or "").strip()
                    if txt:
                        l_list.append(txt)
                elif isinstance(l_item, str):
                    l_list.append(l_item)
            lt = "".join(l_list)
        if lt:
            lyric_lines = [lt]

    return chord_text, note_text, lyric_lines


def _normalize_spacing(text: str, max_consecutive_spaces: int = 2) -> str:
    """壓縮過長連續空白，保留適當間距但防止在 Word 表格內擠到下一行。"""
    if not text:
        return ""
    s = text.replace("\t", " ")
    if max_consecutive_spaces <= 1:
        return re.sub(r" {2,}", " ", s)
    return re.sub(r" {" + str(max_consecutive_spaces + 1) + r",}", " " * max_consecutive_spaces, s)


def _visual_width_units(text: str) -> float:
    """估算字串在單行顯示所需的視覺寬度單位（半形約 1.0，中文字約 2.0）。"""
    s = re.sub(r"[\u0307\u0323]", "", text or "")
    w = 0.0
    for ch in s:
        code = ord(ch)
        if code == 32:
            w += 0.55
        elif code > 255:
            w += 1.95
        else:
            w += 0.9
    return w


def _parse_jianpu_runs(note_text: str) -> List[Tuple[str, int]]:
    """
    將簡譜音符字串解析為 (顯示文字, 底線數 0/1/2) 的序列：
    - 高八度 ' 轉為數字上方點 (\u0307)
    - 低八度 , 轉為數字下方點 (\u0323)
    - _ 轉為單底線 (1)、= 轉為雙底線 (2)
    - ^ 轉為連音線 ⌒、~ 轉為圓滑線 ⌣、(3) 轉為 ³
    """
    s = _normalize_spacing(note_text.strip(), max_consecutive_spaces=2)
    if not s:
        return [("0   -   -   -", 0)]

    tokens = re.split(r"(\s+)", s)
    runs: List[Tuple[str, int]] = []
    for tok in tokens:
        if not tok:
            continue
        if tok.isspace():
            runs.append((tok, 0))
            continue

        # 處理括號三連音標記
        t = tok.replace("(3)", "³")
        has_tie = "^" in t
        has_slur = "~" in t
        t = t.replace("^", "").replace("~", "")

        # 計算減時線底線數
        eq_count = t.count("=")
        us_count = t.count("_")
        uline = 2 if eq_count > 0 else (1 if us_count >= 2 else (1 if us_count == 1 else 0))
        if us_count >= 2 and eq_count == 0:
            uline = 2
        t = t.replace("_", "").replace("=", "")

        # 計算高低八度點與附點
        up_dots = t.count("'")
        down_dots = t.count(",")
        aug_dot = "." in t
        t = t.replace("'", "").replace(",", "").replace(".", "")

        # 組合升降記號與主音符數字
        m = re.match(r"^([#b♯♭]?)([0-7\-xX])(.*)$", t)
        if m:
            acc, digit, rest = m.group(1), m.group(2), m.group(3)
            acc = acc.replace("♯", "#").replace("♭", "b")
            comb = digit
            if up_dots > 0:
                comb += "\u0307" * min(up_dots, 2)
            if down_dots > 0:
                comb += "\u0323" * min(down_dots, 2)
            disp = f"{acc}{comb}{'.' if aug_dot else ''}{rest}"
        else:
            disp = t + ("." if aug_dot else "")

        if has_tie:
            disp += "⌒"
        if has_slur:
            disp += "⌣"

        runs.append((disp, uline))

    return runs


def export_score_to_docx(doc_data: Dict[str, Any]) -> bytes:
    """接收 ScoreDocument dict，回傳 .docx 二進位 bytes。"""
    word_doc = docx.Document()

    # 設定 A4 窄邊界（左右 1.0 cm，讓每行 4 小節有 19.0 cm 寬裕空間不折行）
    for section in word_doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.2)
        section.bottom_margin = Cm(1.2)
        section.left_margin = Cm(1.0)
        section.right_margin = Cm(1.0)

    title = str(doc_data.get("title") or "未命名樂譜").strip()
    subtitle = str(doc_data.get("subtitle") or "").strip()
    key_sig = str(doc_data.get("key") or "C").strip()
    time_sig = str(doc_data.get("timeSignature") or "4/4").strip()
    tempo = doc_data.get("tempo") or 80
    lyricist = str(doc_data.get("lyricist") or doc_data.get("author") or "").strip()
    composer = str(doc_data.get("composer") or doc_data.get("arranger") or "").strip()
    title_scale = float(doc_data.get("titleFontScale") or 1.0)

    # 1. 標題（置中）
    title_p = word_doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_run = title_p.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(max(14, min(28, round(18 * title_scale))))
    title_run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    _set_east_asia_font(title_run)

    if subtitle:
        sub_p = word_doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_p.paragraph_format.space_before = Pt(0)
        sub_p.paragraph_format.space_after = Pt(2)
        sub_run = sub_p.add_run(subtitle)
        sub_run.font.size = Pt(10.5)
        sub_run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
        _set_east_asia_font(sub_run)

    # 2. 標頭資訊列：左方「1=C  4/4  ♩=80」，右方「作詞  作曲」（無框線 2 欄表格）
    meta_tbl = word_doc.add_table(rows=1, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_tbl.autofit = False
    left_cell, right_cell = meta_tbl.rows[0].cells[0], meta_tbl.rows[0].cells[1]
    _set_cell_width_cm(left_cell, 9.5)
    _set_cell_width_cm(right_cell, 9.5)
    _set_cell_borders(left_cell)
    _set_cell_borders(right_cell)
    _set_cell_margins(left_cell, top=20, bottom=60, left=20, right=20)
    _set_cell_margins(right_cell, top=20, bottom=60, left=20, right=20)

    lp = left_cell.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lp.paragraph_format.space_before = Pt(0)
    lp.paragraph_format.space_after = Pt(4)
    lrun = lp.add_run(f"1 = {key_sig}    {time_sig}    ♩ = {tempo}")
    lrun.bold = True
    lrun.font.size = Pt(10.5)
    lrun.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    _set_east_asia_font(lrun)

    rp = right_cell.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.paragraph_format.space_before = Pt(0)
    rp.paragraph_format.space_after = Pt(4)
    credits_parts = []
    if lyricist:
        credits_parts.append(f"作詞：{lyricist}")
    if composer:
        credits_parts.append(f"作曲：{composer}")
    if not credits_parts:
        credits_parts.append("作詞：______    作曲：______")
    rrun = rp.add_run("    ".join(credits_parts))
    rrun.font.size = Pt(10)
    rrun.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    _set_east_asia_font(rrun)

    measures = doc_data.get("measures") or []
    if not isinstance(measures, list) or len(measures) == 0:
        empty_p = word_doc.add_paragraph("（尚無小節內容）")
        empty_run = empty_p.runs[0]
        _set_east_asia_font(empty_run)
    else:
        cols_per_row = 4
        total_width_cm = 19.0
        barline_spec = {"val": "single", "sz": 8, "color": "333333"}

        for row_start in range(0, len(measures), cols_per_row):
            row_measures = measures[row_start : row_start + cols_per_row]

            # 計算本行 4 小節各自的內容寬度權重，動態分配欄寬，讓字多的小節自動獲得更寬空間
            weights: List[float] = []
            parsed_cells = []
            for col_idx in range(cols_per_row):
                if col_idx < len(row_measures) and isinstance(row_measures[col_idx], dict):
                    m = row_measures[col_idx]
                    ch_t, nt_t, lyr_lines = _extract_measure_texts(m)
                    nt_clean = _normalize_spacing(nt_t, 2)
                    lyr_clean = [_normalize_spacing(line, 1) for line in lyr_lines if line.strip()]
                    max_u = max(
                        [_visual_width_units(nt_clean), _visual_width_units(ch_t)]
                        + [_visual_width_units(line) for line in lyr_clean]
                        + [8.0]
                    )
                    weights.append(max(10.0, max_u + 2.5))
                    parsed_cells.append((m, ch_t, nt_clean, lyr_clean))
                else:
                    weights.append(10.0)
                    parsed_cells.append(None)

            sum_w = sum(weights) or 40.0
            col_widths = [max(3.2, total_width_cm * (w / sum_w)) for w in weights]
            scale_fix = total_width_cm / sum(col_widths)
            col_widths = [w * scale_fix for w in col_widths]

            table = word_doc.add_table(rows=1, cols=cols_per_row)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            row_cells = table.rows[0].cells

            for col_idx in range(cols_per_row):
                cell = row_cells[col_idx]
                w_cm = col_widths[col_idx]
                _set_cell_width_cm(cell, w_cm)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                _set_cell_margins(cell, top=40, bottom=50, left=70, right=70)

                item = parsed_cells[col_idx]
                if item is None:
                    _set_cell_borders(cell, top=None, bottom=None, left=None, right=None)
                    continue

                m, chord_text, note_text, lyric_lines = item
                m_idx = row_start + col_idx

                # 與 Score Studio 一致：共用單一小節垂直線（上下無外框線）
                _set_cell_borders(
                    cell,
                    top=None,
                    bottom=None,
                    left=barline_spec,
                    right=barline_spec,
                )

                # 依該欄字數與分配到的欄寬自動微調字級，確保單行不換行
                max_units = weights[col_idx] - 2.5
                avail_chars = (w_cm - 0.3) * 4.5
                font_scale = min(1.0, max(0.78, avail_chars / max(1.0, max_units)))
                note_pt = round(11.0 * font_scale, 1)
                lyric_pt = round(10.0 * font_scale, 1)

                # 第一行：小節編號 + 段落標籤
                sec_raw = str(m.get("section") or "").strip()
                sec_label = SECTION_LABELS.get(sec_raw.lower(), sec_raw) if sec_raw else ""
                p_head = cell.paragraphs[0]
                p_head.paragraph_format.space_before = Pt(0)
                p_head.paragraph_format.space_after = Pt(1)
                p_head.paragraph_format.line_spacing = 1.0

                r_num = p_head.add_run(f"{m_idx + 1}")
                r_num.font.size = Pt(8)
                r_num.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
                _set_east_asia_font(r_num)

                if sec_label:
                    r_sec = p_head.add_run(f"  [{sec_label}]")
                    r_sec.bold = True
                    r_sec.font.size = Pt(8.5)
                    r_sec.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
                    _set_east_asia_font(r_sec)

                # 第二行：和弦行（始終保留高度讓各小節音符水平對齊）
                p_chord = cell.add_paragraph()
                p_chord.paragraph_format.space_before = Pt(0)
                p_chord.paragraph_format.space_after = Pt(1)
                p_chord.paragraph_format.line_spacing = 1.05
                r_chord = p_chord.add_run(chord_text if chord_text else " ")
                r_chord.bold = True
                r_chord.font.size = Pt(max(8.5, round(9.5 * font_scale, 1)))
                r_chord.font.color.rgb = RGBColor(0x19, 0x76, 0xD2)
                _set_east_asia_font(r_chord, ascii_font="Consolas")

                # 第三行：簡譜音符行（含高低八度點與減時線底線）
                p_note = cell.add_paragraph()
                p_note.paragraph_format.space_before = Pt(1)
                p_note.paragraph_format.space_after = Pt(2)
                p_note.paragraph_format.line_spacing = 1.15
                for seg_text, uline in _parse_jianpu_runs(note_text):
                    r_note = p_note.add_run(seg_text)
                    r_note.bold = True
                    r_note.font.size = Pt(note_pt)
                    r_note.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
                    if uline == 1:
                        r_note.underline = WD_UNDERLINE.SINGLE
                    elif uline >= 2:
                        r_note.underline = WD_UNDERLINE.DOUBLE
                    _set_east_asia_font(r_note, ascii_font="Consolas")

                # 第四行起：歌詞行
                if lyric_lines:
                    for line in lyric_lines:
                        p_lyr = cell.add_paragraph()
                        p_lyr.paragraph_format.space_before = Pt(0)
                        p_lyr.paragraph_format.space_after = Pt(1)
                        p_lyr.paragraph_format.line_spacing = 1.1
                        r_lyr = p_lyr.add_run(line)
                        r_lyr.font.size = Pt(lyric_pt)
                        r_lyr.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
                        _set_east_asia_font(r_lyr)
                else:
                    p_lyr = cell.add_paragraph()
                    p_lyr.paragraph_format.space_before = Pt(0)
                    p_lyr.paragraph_format.space_after = Pt(1)
                    r_lyr = p_lyr.add_run(" ")
                    r_lyr.font.size = Pt(lyric_pt)

            # 每行 4 小節之間留白間距
            spacer = word_doc.add_paragraph()
            spacer.paragraph_format.space_before = Pt(0)
            spacer.paragraph_format.space_after = Pt(4)
            spacer_run = spacer.add_run("")
            spacer_run.font.size = Pt(4)

    buf = io.BytesIO()
    word_doc.save(buf)
    return buf.getvalue()


export_score_to_word = export_score_to_docx

